"""Create a machine-readable inventory of N100 Excel source files."""

from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def serialise(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    return value


def value_type(value: Any) -> str:
    if value is None or (isinstance(value, str) and not value.strip()):
        return "blank"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime, date)):
        return "date"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    return "text"


def detect_header(rows: list[list[Any]], scan_rows: int = 15) -> int:
    """Return the zero-based row index with the most populated header cells."""
    candidates = rows[:scan_rows]
    return max(
        range(len(candidates)),
        key=lambda idx: (
            sum(value_type(v) != "blank" for v in candidates[idx]),
            sum(value_type(v) == "text" for v in candidates[idx]),
        ),
    )


def inspect_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    file_result: dict[str, Any] = {
        "file": path.name,
        "size_bytes": path.stat().st_size,
        "sheets": [],
    }
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        while rows and all(value_type(v) == "blank" for v in rows[-1]):
            rows.pop()
        if not rows:
            file_result["sheets"].append({"sheet": sheet.title, "rows": 0, "columns": 0})
            continue

        last_column = max(
            max((i + 1 for i, value in enumerate(row) if value_type(value) != "blank"), default=0)
            for row in rows
        )
        rows = [row[:last_column] + [None] * max(0, last_column - len(row)) for row in rows]
        header_index = detect_header(rows)
        headers = [str(v).strip() if v is not None else None for v in rows[header_index]]
        data = [row for row in rows[header_index + 1:] if any(value_type(v) != "blank" for v in row)]

        columns = []
        for column_index, header in enumerate(headers):
            values = [row[column_index] for row in data]
            nonblank = [v for v in values if value_type(v) != "blank"]
            columns.append(
                {
                    "name": header,
                    "nonblank": len(nonblank),
                    "blank": len(values) - len(nonblank),
                    "unique": len({json.dumps(serialise(v), default=str) for v in nonblank}),
                    "types": dict(Counter(value_type(v) for v in nonblank)),
                    "sample": [serialise(v) for v in nonblank[:3]],
                }
            )

        file_result["sheets"].append(
            {
                "sheet": sheet.title,
                "rows_including_header": len(rows),
                "data_rows": len(data),
                "columns": len(headers),
                "header_row": header_index + 1,
                "headers": headers,
                "column_profile": columns,
            }
        )
    workbook.close()
    return file_result


def build_inventory(input_dir: Path, additional_dir: Path | None = None) -> dict[str, Any]:
    directories = [input_dir]
    if additional_dir is not None:
        directories.append(additional_dir)
    files = sorted(path for directory in directories for path in directory.glob("*.xlsx"))
    return {
        "input_directories": [str(directory) for directory in directories],
        "workbook_count": len(files),
        "workbooks": [inspect_workbook(path) for path in files],
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--additional-dir", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    inventory = build_inventory(args.input_dir, args.additional_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(inventory, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Inventoried {inventory['workbook_count']} workbooks -> {args.output}")


if __name__ == "__main__":
    main()
