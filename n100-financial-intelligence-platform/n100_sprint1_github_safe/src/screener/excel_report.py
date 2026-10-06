"""Generate the Sprint 3 Day 17 six-sheet screener workbook."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any, Mapping, Sequence

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.screener.composite_score import build_scored_snapshot
from src.screener.engine import DEFAULT_CONFIG_PATH, DEFAULT_DATABASE_PATH
from src.screener.presets import (
    EXPECTED_PRESET_KEYS,
    apply_preset,
    evaluate_rule,
    load_preset_definitions,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DEFAULT_OUTPUT = PROJECT_ROOT / "output" / "screener_output.xlsx"

KPI_COLUMNS = [
    "sprint3_composite_score",
    "sector_relative_composite_score",
    "return_on_equity_pct",
    "return_on_capital_employed_pct",
    "net_profit_margin_pct",
    "operating_profit_margin_pct",
    "free_cash_flow_cr",
    "fcf_cagr_5yr",
    "cfo_pat_ratio_5yr",
    "revenue_cagr_3yr",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "eps_cagr_5yr",
    "debt_to_equity",
    "interest_coverage",
    "pe_ratio",
    "pb_ratio",
    "dividend_yield_pct",
    "dividend_payout_ratio_pct",
    "sales",
]

AUDIT_COLUMNS = [
    "legacy_composite_quality_score",
    "global_score_coverage_pct",
    "sector_score_coverage_pct",
    "previous_debt_to_equity",
    "fcf_cagr_5yr_flag",
]

IDENTIFIER_COLUMNS = [
    "preset_rank",
    "company_id",
    "company_name",
    "broad_sector",
    "financial_year",
]

SHEET_NAMES = {
    "quality_compounder": "Quality Compounder",
    "value_pick": "Value Pick",
    "growth_accelerator": "Growth Accelerator",
    "dividend_champion": "Dividend Champion",
    "debt_free_blue_chip": "Debt-Free Blue Chip",
    "turnaround_watch": "Turnaround Watch",
}

HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
PASS_FILL = PatternFill("solid", fgColor="C6EFCE")
FAIL_FILL = PatternFill("solid", fgColor="FFC7CE")


def _presentation_frame(result: pd.DataFrame) -> pd.DataFrame:
    ordered = result.sort_values(
        ["sprint3_composite_score", "company_id"],
        ascending=[False, True],
        na_position="last",
        kind="mergesort",
    ).reset_index(drop=True)
    ordered.insert(0, "preset_rank", range(1, len(ordered) + 1))
    return ordered[IDENTIFIER_COLUMNS + AUDIT_COLUMNS + KPI_COLUMNS]


def build_preset_frames(
    scored_snapshot: pd.DataFrame,
    definitions: Mapping[str, Mapping[str, Any]],
) -> dict[str, pd.DataFrame]:
    """Return the six non-empty, new-score-sorted result frames."""

    frames: dict[str, pd.DataFrame] = {}
    for key in EXPECTED_PRESET_KEYS:
        result, _ = apply_preset(scored_snapshot, definitions[key])
        if result.empty:
            raise ValueError(f"Official preset '{key}' produced an empty worksheet")
        frames[key] = _presentation_frame(result)
    return frames


def _rule_display_column(rule: Mapping[str, Any]) -> str:
    if rule.get("special") == "debt_to_equity_declining":
        return "debt_to_equity"
    return str(rule["column"])


def _apply_styles(
    workbook_path: Path,
    source_frames: Mapping[str, pd.DataFrame],
    definitions: Mapping[str, Mapping[str, Any]],
) -> None:
    workbook = load_workbook(workbook_path)
    for key in EXPECTED_PRESET_KEYS:
        worksheet = workbook[SHEET_NAMES[key]]
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.showGridLines = False
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = worksheet.PAPERSIZE_A3
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.print_title_rows = "1:1"
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        worksheet.row_dimensions[1].height = 32

        headers = {cell.value: cell.column for cell in worksheet[1]}
        source = source_frames[key].reset_index(drop=True)
        for rule in definitions[key]["rules"]:
            display_column = _rule_display_column(rule)
            if display_column not in headers:
                raise ValueError(
                    f"Workbook is missing threshold column '{display_column}'"
                )
            column_number = headers[display_column]
            worksheet.cell(1, column_number).comment = Comment(
                str(rule["label"]), "N100 Sprint 3"
            )
            passed = evaluate_rule(source, rule)
            for row_offset, rule_passed in enumerate(passed.tolist(), start=2):
                worksheet.cell(row_offset, column_number).fill = (
                    PASS_FILL if rule_passed else FAIL_FILL
                )

        for column_number, cell in enumerate(worksheet[1], start=1):
            header = str(cell.value)
            values = [str(worksheet.cell(row, column_number).value or "") for row in range(2, worksheet.max_row + 1)]
            width = max([len(header), *(len(value) for value in values)] or [len(header)]) + 2
            worksheet.column_dimensions[get_column_letter(column_number)].width = min(max(width, 11), 34)

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.00"
                cell.alignment = Alignment(vertical="top")
    workbook.save(workbook_path)


def write_screener_workbook(
    scored_snapshot: pd.DataFrame,
    definitions: Mapping[str, Mapping[str, Any]],
    output_path: str | Path = DEFAULT_OUTPUT,
) -> dict[str, pd.DataFrame]:
    """Write exactly six formatted preset sheets and return source frames."""

    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    frames = build_preset_frames(scored_snapshot, definitions)
    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        for key in EXPECTED_PRESET_KEYS:
            frames[key].to_excel(writer, sheet_name=SHEET_NAMES[key], index=False)
    _apply_styles(destination, frames, definitions)
    return frames


def validate_workbook(
    workbook_path: str | Path,
) -> pd.DataFrame:
    """Validate sheet count, rows, KPI presence and presentation controls."""

    workbook = load_workbook(workbook_path, data_only=False)
    if workbook.sheetnames != [SHEET_NAMES[key] for key in EXPECTED_PRESET_KEYS]:
        raise ValueError("Workbook does not contain the six official sheets in order")
    rows: list[dict[str, Any]] = []
    for worksheet in workbook.worksheets:
        headers = [cell.value for cell in worksheet[1]]
        missing_kpis = sorted(set(KPI_COLUMNS).difference(headers))
        data_rows = worksheet.max_row - 1
        status = (
            "PASS"
            if data_rows > 0
            and not missing_kpis
            and worksheet.freeze_panes == "A2"
            and bool(worksheet.auto_filter.ref)
            else "FAIL"
        )
        rows.append(
            {
                "sheet_name": worksheet.title,
                "data_rows": data_rows,
                "kpi_columns": len([column for column in KPI_COLUMNS if column in headers]),
                "missing_kpis": " | ".join(missing_kpis),
                "freeze_panes": str(worksheet.freeze_panes),
                "auto_filter": worksheet.auto_filter.ref,
                "status": status,
            }
        )
    return pd.DataFrame(rows)


def run_day17(
    database_path: str | Path = DEFAULT_DATABASE_PATH,
    *,
    config_path: str | Path = DEFAULT_CONFIG_PATH,
    output_path: str | Path = DEFAULT_OUTPUT,
) -> dict[str, Any]:
    """Execute scoring, exports and workbook validation."""

    destination = Path(output_path)
    scored, bounds = build_scored_snapshot(database_path)
    definitions = load_preset_definitions(config_path)
    frames = write_screener_workbook(scored, definitions, destination)
    validation = validate_workbook(destination)
    output_dir = destination.parent
    scored.to_csv(output_dir / "composite_score_audit.csv", index=False)
    bounds.to_csv(output_dir / "composite_score_bounds.csv", index=False)
    validation.to_csv(output_dir / "day17_workbook_validation.csv", index=False)
    return {
        "scored_snapshot": scored,
        "bounds": bounds,
        "preset_frames": frames,
        "validation": validation,
        "output_path": destination,
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--database", type=Path, default=DEFAULT_DATABASE_PATH)
    parser.add_argument("--config", type=Path, default=DEFAULT_CONFIG_PATH)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args(argv)
    result = run_day17(
        args.database, config_path=args.config, output_path=args.output
    )
    scored = result["scored_snapshot"]
    validation = result["validation"]
    print(
        json.dumps(
            {
                "companies_scored": int(len(scored)),
                "sheets": int(len(validation)),
                "all_sheets_pass": bool(validation["status"].eq("PASS").all()),
                "sheet_rows": dict(zip(validation["sheet_name"], validation["data_rows"])),
                "output": str(args.output),
            },
            indent=2,
            sort_keys=True,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
