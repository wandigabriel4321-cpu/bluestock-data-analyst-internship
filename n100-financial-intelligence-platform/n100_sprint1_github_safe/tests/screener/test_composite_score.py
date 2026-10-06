"""Tests for Sprint 3 Day 17 composite scoring and Excel output."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

from src.analytics.cagr import FLAG_OK, FLAG_TURNAROUND
from src.screener.composite_score import (
    COMPONENTS,
    ScoreComponent,
    calculate_fcf_cagr_5yr,
    compute_composite_scores,
    validate_component_weights,
    winsorized_minmax,
)
from src.screener.excel_report import KPI_COLUMNS, SHEET_NAMES, _apply_styles, run_day17
from src.screener.presets import load_preset_definitions


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATABASE_PATH = PROJECT_ROOT / "db" / "nifty100.db"
CONFIG_PATH = PROJECT_ROOT / "config" / "screener_config.yaml"


def score_snapshot() -> pd.DataFrame:
    rows = []
    for index, company_id in enumerate(("AAA", "BBB", "CCC"), start=1):
        rows.append(
            {
                "company_id": company_id,
                "company_name": company_id,
                "broad_sector": "Technology" if company_id != "CCC" else "Industrials",
                "financial_year": "2024-03",
                "return_on_equity_pct": 10 * index,
                "return_on_capital_employed_pct": 8 * index,
                "net_profit_margin_pct": 5 * index,
                "free_cash_flow_cr": 100 * index,
                "cfo_pat_ratio_5yr": 0.5 * index,
                "revenue_cagr_5yr": 5 * index,
                "pat_cagr_5yr": 4 * index,
                "debt_to_equity": 2.0 / index,
                "interest_coverage": 2 * index,
                "icr_label": None,
                "composite_quality_score": 20 * index,
            }
        )
    return pd.DataFrame(rows)


def fcf_history() -> pd.DataFrame:
    rows = []
    for index, company_id in enumerate(("AAA", "BBB", "CCC"), start=1):
        rows.extend(
            [
                {"company_id": company_id, "year": "2019-03", "free_cash_flow_cr": 50 * index},
                {"company_id": company_id, "year": "2024-03", "free_cash_flow_cr": 100 * index},
            ]
        )
    return pd.DataFrame(rows)


class CompositeFormulaTests(unittest.TestCase):
    def test_01_formula_has_ten_components_and_weights_sum_to_one(self):
        validate_component_weights()
        self.assertEqual(len(COMPONENTS), 10)
        self.assertAlmostEqual(sum(item.weight for item in COMPONENTS.values()), 1.0)

    def test_02_invalid_weight_total_is_rejected(self):
        invalid = dict(COMPONENTS)
        invalid["roe"] = ScoreComponent("return_on_equity_pct", 0.99)
        with self.assertRaisesRegex(ValueError, "sum to 1.0"):
            validate_component_weights(invalid)

    def test_03_p10_p90_winsorisation_caps_extreme_values(self):
        values = pd.Series(range(0, 101, 10), dtype=float)
        score, lower, upper = winsorized_minmax(values)
        self.assertEqual(lower, 10)
        self.assertEqual(upper, 90)
        self.assertEqual(score.iloc[0], 0)
        self.assertEqual(score.iloc[-1], 100)

    def test_04_inverse_metric_rewards_lower_values(self):
        score, _, _ = winsorized_minmax(pd.Series([1.0, 2.0, 3.0]), inverse=True)
        self.assertGreater(score.iloc[0], score.iloc[-1])

    def test_05_constant_distribution_is_neutral_and_missing_stays_missing(self):
        score, lower, upper = winsorized_minmax(pd.Series([5.0, 5.0, np.nan]))
        self.assertEqual((lower, upper), (5.0, 5.0))
        self.assertEqual(score.iloc[0], 50)
        self.assertTrue(pd.isna(score.iloc[2]))

    def test_06_fcf_cagr_uses_exact_five_year_window(self):
        result = calculate_fcf_cagr_5yr(
            [
                {"year": "2019-03", "free_cash_flow_cr": 100},
                {"year": "2024-03", "free_cash_flow_cr": 200},
            ],
            end_year="2024-03",
        )
        self.assertEqual(result.flag, FLAG_OK)
        self.assertAlmostEqual(result.value, (2 ** (1 / 5) - 1) * 100)

    def test_07_negative_to_positive_fcf_is_flagged_not_fractionally_powered(self):
        result = calculate_fcf_cagr_5yr(
            [
                {"year": "2019-03", "free_cash_flow_cr": -100},
                {"year": "2024-03", "free_cash_flow_cr": 200},
            ],
            end_year="2024-03",
        )
        self.assertEqual(result.flag, FLAG_TURNAROUND)
        self.assertIsNone(result.value)

    def test_08_new_score_preserves_legacy_score_in_separate_column(self):
        scored, _ = compute_composite_scores(score_snapshot(), fcf_history())
        self.assertIn("legacy_composite_quality_score", scored.columns)
        self.assertIn("sprint3_composite_score", scored.columns)
        self.assertEqual(scored["legacy_composite_quality_score"].tolist(), [20, 40, 60])

    def test_09_complete_components_produce_full_coverage(self):
        scored, _ = compute_composite_scores(score_snapshot(), fcf_history())
        self.assertTrue(scored["global_score_coverage_pct"].eq(100).all())
        self.assertTrue(scored["sprint3_composite_score"].between(0, 100).all())

    def test_10_missing_component_is_not_imputed_as_zero(self):
        snapshot = score_snapshot()
        snapshot.loc[snapshot["company_id"].eq("AAA"), "return_on_equity_pct"] = None
        scored, _ = compute_composite_scores(snapshot, fcf_history())
        row = scored.loc[scored["company_id"].eq("AAA")].iloc[0]
        self.assertEqual(row["global_score_coverage_pct"], 85)
        self.assertTrue(0 <= row["sprint3_composite_score"] <= 100)

    def test_11_sector_relative_score_is_computed_separately(self):
        scored, bounds = compute_composite_scores(score_snapshot(), fcf_history())
        self.assertTrue(scored["sector_relative_composite_score"].notna().all())
        self.assertTrue(bounds["scope"].str.startswith("SECTOR:").any())


class ScreenerWorkbookTests(unittest.TestCase):
    def test_12_workbook_has_exactly_six_nonempty_sheets_and_twenty_kpis(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "screener_output.xlsx"
            result = run_day17(
                DATABASE_PATH, config_path=CONFIG_PATH, output_path=path
            )
            self.assertTrue(result["validation"]["status"].eq("PASS").all())
            workbook = load_workbook(path)
            self.assertEqual(workbook.sheetnames, list(SHEET_NAMES.values()))
            self.assertEqual(len(KPI_COLUMNS), 20)
            for worksheet in workbook.worksheets:
                self.assertGreater(worksheet.max_row, 1)

    def test_13_workbook_has_filters_frozen_headers_and_threshold_colours(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "screener_output.xlsx"
            run_day17(DATABASE_PATH, config_path=CONFIG_PATH, output_path=path)
            workbook = load_workbook(path)
            worksheet = workbook["Quality Compounder"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            self.assertEqual(worksheet.freeze_panes, "A2")
            self.assertTrue(worksheet.auto_filter.ref)
            roe_cell = worksheet.cell(2, headers["return_on_equity_pct"])
            self.assertEqual(roe_cell.fill.fgColor.rgb, "00C6EFCE")
            self.assertEqual(worksheet.cell(1, 1).fill.fgColor.rgb, "0017365D")

    def test_14_threshold_failure_is_coloured_red(self):
        with tempfile.TemporaryDirectory() as directory:
            original = Path(directory) / "original.xlsx"
            result = run_day17(
                DATABASE_PATH, config_path=CONFIG_PATH, output_path=original
            )
            frames = {key: frame.copy() for key, frame in result["preset_frames"].items()}
            frames["quality_compounder"].loc[0, "return_on_equity_pct"] = 0
            path = Path(directory) / "red_check.xlsx"
            with pd.ExcelWriter(path, engine="openpyxl") as writer:
                for key, frame in frames.items():
                    frame.to_excel(writer, sheet_name=SHEET_NAMES[key], index=False)
            definitions = load_preset_definitions(CONFIG_PATH)
            _apply_styles(path, frames, definitions)
            workbook = load_workbook(path)
            worksheet = workbook["Quality Compounder"]
            headers = {cell.value: cell.column for cell in worksheet[1]}
            roe_cell = worksheet.cell(2, headers["return_on_equity_pct"])
            self.assertEqual(roe_cell.fill.fgColor.rgb, "00FFC7CE")


if __name__ == "__main__":
    unittest.main()
