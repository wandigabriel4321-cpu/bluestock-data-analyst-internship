"""Tests for Sprint 3 Day 20 peer-comparison workbook."""

from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.analytics.peer_report import (
    BENCHMARK_FILL,
    HIGH_FILL,
    KPI_COLUMNS,
    LOW_FILL,
    MEDIAN_FILL,
    MEDIAN_ID,
    MID_FILL,
    OFFICIAL_GROUP_ORDER,
    OUTPUT_COLUMNS,
    PERCENTILE_COLUMNS,
    build_group_frames,
    build_peer_report_rows,
    percentile_fill,
    validate_peer_workbook,
    write_peer_workbook,
)


PROJECT_ROOT = Path(__file__).resolve().parents[2]
DATABASE_PATH = PROJECT_ROOT / "db" / "nifty100.db"


class PeerReportStructureTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows = build_peer_report_rows(DATABASE_PATH)
        cls.frames = build_group_frames(cls.rows)
        cls.benchmarks = {
            str(group): str(frame.loc[frame["is_benchmark"].eq(1), "company_id"].iloc[0])
            for group, frame in cls.rows.groupby("peer_group_name", sort=False)
        }

    def test_01_report_has_exactly_twenty_kpis(self):
        self.assertEqual(len(KPI_COLUMNS), 20)

    def test_02_report_has_exactly_ten_official_percentiles(self):
        self.assertEqual(len(PERCENTILE_COLUMNS), 10)
        self.assertEqual(len(set(PERCENTILE_COLUMNS.values())), 10)

    def test_03_output_has_two_identifiers_and_thirty_metric_columns(self):
        self.assertEqual(len(OUTPUT_COLUMNS), 32)
        self.assertEqual(OUTPUT_COLUMNS[:2], ("company_id", "company_name"))

    def test_04_real_rows_cover_56_unique_companies_and_11_groups(self):
        self.assertEqual(len(self.rows), 56)
        self.assertEqual(self.rows["company_id"].nunique(), 56)
        self.assertEqual(self.rows["peer_group_name"].nunique(), 11)

    def test_05_every_group_has_exactly_one_benchmark(self):
        counts = self.rows.groupby("peer_group_name")["is_benchmark"].sum()
        self.assertTrue(counts.eq(1).all())

    def test_06_frames_preserve_official_order_and_end_with_median(self):
        self.assertEqual(tuple(self.frames), OFFICIAL_GROUP_ORDER)
        for frame in self.frames.values():
            self.assertEqual(frame.iloc[-1]["company_id"], MEDIAN_ID)

    def test_07_median_values_are_recomputed_from_company_rows(self):
        for frame in self.frames.values():
            companies = frame.iloc[:-1]
            median = frame.iloc[-1]
            for column in (*KPI_COLUMNS, *PERCENTILE_COLUMNS.values()):
                values = pd.to_numeric(companies[column], errors="coerce")
                if values.notna().any():
                    self.assertAlmostEqual(float(median[column]), float(values.median()))

    def test_08_available_percentiles_are_bounded_zero_to_one_hundred(self):
        values = self.rows[list(PERCENTILE_COLUMNS.values())].apply(pd.to_numeric, errors="coerce").stack()
        self.assertTrue(values.between(0, 100, inclusive="both").all())

    def test_09_official_raw_value_matches_latest_percentile_source(self):
        for metric, percentile_column in PERCENTILE_COLUMNS.items():
            available = self.rows[percentile_column].notna()
            self.assertTrue(self.rows.loc[available, metric].notna().all())


class PeerReportStyleTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rows = build_peer_report_rows(DATABASE_PATH)
        cls.frames = build_group_frames(cls.rows)
        cls.benchmarks = {
            str(group): str(frame.loc[frame["is_benchmark"].eq(1), "company_id"].iloc[0])
            for group, frame in cls.rows.groupby("peer_group_name", sort=False)
        }

    def test_10_percentile_boundary_colours_follow_official_precedence(self):
        self.assertEqual(percentile_fill(75).fgColor.rgb, HIGH_FILL.fgColor.rgb)
        self.assertEqual(percentile_fill(50).fgColor.rgb, MID_FILL.fgColor.rgb)
        self.assertEqual(percentile_fill(25).fgColor.rgb, LOW_FILL.fgColor.rgb)

    def test_11_workbook_has_11_sheets_filters_freeze_and_valid_audit(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "peer_comparison.xlsx"
            write_peer_workbook(self.frames, self.benchmarks, path)
            audit = validate_peer_workbook(path, self.benchmarks)
            workbook = load_workbook(path)
            self.assertEqual(tuple(workbook.sheetnames), OFFICIAL_GROUP_ORDER)
            self.assertTrue(audit["status"].eq("PASS").all())
            for sheet in workbook.worksheets:
                self.assertEqual(sheet.freeze_panes, "C2")
                self.assertTrue(sheet.auto_filter.ref)

    def test_12_benchmark_is_gold_and_median_has_distinct_style(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "peer_comparison.xlsx"
            write_peer_workbook(self.frames, self.benchmarks, path)
            workbook = load_workbook(path)
            sheet = workbook["IT Services"]
            ids = {sheet.cell(row, 1).value: row for row in range(2, sheet.max_row + 1)}
            benchmark_row = ids[self.benchmarks["IT Services"]]
            self.assertEqual(sheet.cell(benchmark_row, 1).fill.fgColor.rgb, BENCHMARK_FILL.fgColor.rgb)
            self.assertEqual(sheet.cell(sheet.max_row, 1).fill.fgColor.rgb, MEDIAN_FILL.fgColor.rgb)

    def test_13_percentage_currency_ratio_and_percentile_formats_are_distinct(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "peer_comparison.xlsx"
            write_peer_workbook(self.frames, self.benchmarks, path)
            sheet = load_workbook(path)["FMCG"]
            headers = {cell.value: cell.column for cell in sheet[1]}
            self.assertIn("%", sheet.cell(2, headers["return_on_equity_pct"]).number_format)
            self.assertIn("₹", sheet.cell(2, headers["free_cash_flow_cr"]).number_format)
            self.assertIn("x", sheet.cell(2, headers["debt_to_equity"]).number_format)
            self.assertEqual(sheet.cell(2, headers["return_on_equity_percentile"]).number_format, "0.00")


if __name__ == "__main__":
    unittest.main()
